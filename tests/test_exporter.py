import datetime as dt

import openpyxl
import pandas as pd

from core import exporter
from core.validator import Issue


def test_export_format(tmp_path):
    summary = pd.DataFrame([["一店", 1234.5, 3]], columns=["门店", "金额(求和)", "记录数"])
    detail = pd.DataFrame([[dt.date(2026, 9, 1), "=SUM(A1)", 1000, "a.xlsx", 2],
                           [dt.datetime(2026, 9, 2, 8, 0), "坏\x01字符", 2, "a.xlsx", 3]],
                          columns=["日期", "备注", "数量", "来源文件", "原始行号"]).astype(object)
    issues = [Issue("错误", "a.xlsx", "S1", 3, "金额", "数值无法转换", "说明", "约5元")]
    out = exporter.export_result(tmp_path / "r.xlsx", summary, detail, issues,
                                 plain_columns=["原始行号"])
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["汇总", "明细", "问题清单"]
    for ws in wb:
        assert ws.freeze_panes == "A2"
        assert ws["A1"].font.b
        assert ws["A1"].fill.fgColor.rgb.endswith("DDEBF7")
    ws = wb["明细"]
    assert ws["B2"].value == "=SUM(A1)" and ws["B2"].data_type == "s"   # 文本不能变成公式
    assert ws["B3"].value == "坏字符"
    assert ws["A2"].number_format == "yyyy-mm-dd hh:mm:ss"  # 列中含时间则统一带时间
    assert ws["C2"].number_format == "#,##0"
    assert ws["E2"].number_format == "General"
    assert wb["汇总"]["B2"].number_format == "#,##0.00"
    assert ws.column_dimensions["A"].width >= 10
    assert wb["问题清单"]["A2"].value == "错误"


def test_default_filename():
    name = exporter.default_filename(dt.datetime(2026, 10, 3, 9, 5, 7))
    assert name == "汇总结果_20261003_090507.xlsx"
