"""端到端：用全部测试数据跑完整流程，验证容错和问题定位。"""
import shutil

import openpyxl
import pytest

import cli
from config.task_config import AggSpec, TaskConfig
from core import pipeline
from core.validator import ERROR, T_BAD_DATE, T_BAD_NUMBER, T_DUPLICATE, T_READ_FAIL


@pytest.fixture
def cfg(data_dir, tmp_path):
    return TaskConfig(input_folder=str(data_dir), group_by=["门店"],
                      aggregations=[AggSpec("金额", "求和")], output_dir=str(tmp_path),
                      output_name="out.xlsx")


def _issues(result, type_=None, file=None):
    return [i for i in result.issues.issues
            if (type_ is None or i.type == type_) and (file is None or i.file == file)]


def test_full_run_no_exception(cfg):
    result = pipeline.run(cfg)
    assert result.output_path and result.output_path.exists()
    assert result.files_total == 19
    assert result.rows_detail == 125
    wb = openpyxl.load_workbook(result.output_path)
    assert wb.sheetnames == ["汇总", "明细", "问题清单"]
    assert wb["明细"].max_row == 126


def test_corrupt_file_reported(cfg):
    result = pipeline.run(cfg)
    (issue,) = _issues(result, T_READ_FAIL, "16_损坏文件.xlsx")
    assert issue.severity == ERROR and "无法打开" in issue.message


def test_issues_locate_file_sheet_row(cfg):
    result = pipeline.run(cfg)
    bad = {(i.file, i.sheet, i.row, i.column, str(i.value)) for i in result.issues.issues
           if i.type in (T_BAD_NUMBER, T_BAD_DATE)}
    assert bad == {
        ("08_文本数字.xlsx", "Sheet1", 7, "金额", "约52元"),
        ("09_日期混用.xlsx", "Sheet1", 10, "日期", "2026-13-45"),
        ("09_日期混用.xlsx", "Sheet1", 11, "日期", "下周一"),
        ("10_占位值.xlsx", "Sheet1", 7, "数量", "很多"),
    }
    # 所有行级问题都带文件和 Sheet
    for i in result.issues.issues:
        if i.row is not None:
            assert i.file and i.sheet


def test_duplicates_reported(cfg):
    result = pipeline.run(cfg)
    dups = sorted((i.file, i.row) for i in _issues(result, T_DUPLICATE))
    assert dups == [("11_重复记录.xlsx", 3), ("11_重复记录.xlsx", 6), ("11_重复记录.xlsx", 7)]


def test_dedup_drop_changes_totals(cfg):
    keep = pipeline.run(cfg)
    cfg.dedup_mode = "drop"
    drop = pipeline.run(cfg)
    assert drop.rows_detail == keep.rows_detail - 3


def test_source_columns_trace_back(cfg):
    result = pipeline.run(cfg)
    d = result.detail
    row = d[(d["来源文件"] == "06_标题行和合计行.xlsx")].iloc[0]
    assert row["来源Sheet"] == "Sheet1" and row["原始行号"] == 5
    assert set(d.columns[-3:]) == {"来源文件", "来源Sheet", "原始行号"}


def test_summary_total_matches_detail(cfg):
    cfg.add_count_column = True
    result = pipeline.run(cfg)
    s = result.summary
    total = s[s["门店"] == "总计"]["金额(求和)"].iloc[0]
    assert total == pytest.approx(sum(v for v in result.detail["金额"] if v is not None))
    assert s[s["门店"] == "总计"]["记录数"].iloc[0] == 125


def test_recursive_includes_subfolder(cfg):
    cfg.recursive = True
    result = pipeline.run(cfg)
    assert "五店" in set(result.detail["门店"])


def test_exclude_and_mapping_override(cfg):
    cfg.excluded = ["12_多Sheet.xlsx|*", "01_标准格式.xlsx|销售明细"]
    cfg.column_mapping = {"会员卡号": ""}
    result = pipeline.run(cfg)
    assert "会员卡号" not in result.detail.columns
    assert not set(result.detail["来源文件"]) & {"12_多Sheet.xlsx", "01_标准格式.xlsx"}


def test_strict_mapping_keeps_original_header(cfg):
    cfg.accept_pending = False
    result = pipeline.run(cfg)
    assert "单价/元" in result.detail.columns


def test_missing_input_folder(tmp_path):
    result = pipeline.run(TaskConfig(input_folder=str(tmp_path / "nope"), output_dir=str(tmp_path)))
    assert result.output_path is None and result.issues.count(ERROR) == 1


def test_output_files_are_not_reread(data_dir, tmp_path):
    folder = tmp_path / "in"
    shutil.copytree(data_dir, folder)
    cfg = TaskConfig(input_folder=str(folder))       # 结果默认写在输入文件夹
    first = pipeline.run(cfg)
    second = pipeline.run(cfg)
    assert first.output_path.parent == folder
    assert second.rows_detail == first.rows_detail


def test_config_roundtrip_same_result(cfg, tmp_path):
    cfg.column_mapping = {"单价/元": "单价"}
    cfg.header_rows = {"06_标题行和合计行.xlsx|Sheet1": 4}
    path = cfg.save(tmp_path / "task.json")
    loaded = TaskConfig.load(path)
    assert loaded == cfg
    a, b = pipeline.run(cfg), pipeline.run(loaded)
    assert a.summary.equals(b.summary) and a.detail.equals(b.detail)
    assert [i.to_row() for i in a.issues.sorted()] == [i.to_row() for i in b.issues.sorted()]


def test_cancel(cfg):
    import threading
    ev = threading.Event()
    ev.set()
    with pytest.raises(pipeline.Cancelled):
        pipeline.run(cfg, cancel=ev)


def test_cli_example(data_dir, tmp_path, capsys):
    out = tmp_path / "result.xlsx"
    code = cli.main(["--input", str(data_dir), "--group", "门店", "--sum", "金额",
                     "--output", str(out), "--save-config", str(tmp_path / "c.json")])
    assert code == 0 and out.exists()
    text = capsys.readouterr().out
    assert "明细 125 行" in text
    # 加载保存的配置一键重跑
    out.unlink()
    assert cli.main(["--config", str(tmp_path / "c.json")]) == 0 and out.exists()


def test_cli_show_mapping(data_dir, capsys):
    assert cli.main(["--input", str(data_dir), "--show-mapping"]) == 0
    assert "待确认" in capsys.readouterr().out


def test_cli_learn(data_dir, tmp_path, isolated_home):
    cli.main(["--input", str(data_dir), "--map", "单价/元=单价", "--learn", "--show-mapping"])
    from core.header_mapper import AliasStore, suggest
    assert suggest("单价/元", AliasStore()).status == "自动"


def test_count_column_default_off_and_saved(cfg, tmp_path):
    assert "记录数" not in pipeline.run(cfg, export=False).summary.columns
    cfg.add_count_column = True
    loaded = TaskConfig.load(cfg.save(tmp_path / "c.json"))
    assert loaded.add_count_column is True
    assert "记录数" in pipeline.run(loaded, export=False).summary.columns
    # 旧版配置文件没有这个字段，按默认不附加
    import json
    data = json.loads((tmp_path / "c.json").read_text(encoding="utf-8"))
    data.pop("add_count_column")
    assert TaskConfig.from_dict(data).add_count_column is False


def test_cli_count_column(data_dir, tmp_path):
    import openpyxl
    out = tmp_path / "r.xlsx"
    assert cli.main(["-i", str(data_dir), "-g", "门店", "-s", "金额", "-o", str(out)]) == 0
    assert [c.value for c in openpyxl.load_workbook(out)["汇总"][1]] == ["门店", "金额(求和)"]
    assert cli.main(["-i", str(data_dir), "-g", "门店", "-s", "金额", "--count-column",
                     "-o", str(out)]) == 0
    assert [c.value for c in openpyxl.load_workbook(out)["汇总"][1]] == ["门店", "金额(求和)", "记录数"]


def test_split_into_sheets(cfg):
    cfg.group_by, cfg.pivot_column, cfg.split_by = ["商品"], "门店", "门店"
    result = pipeline.run(cfg)
    wb = openpyxl.load_workbook(result.output_path)
    assert wb.sheetnames == ["汇总", "一店", "二店", "三店", "四店", "明细", "问题清单"]
    assert [c.value for c in wb["汇总"][1]] == ["商品", "一店", "二店", "三店", "四店", "合计"]
    ws = wb["二店"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    detail_amount = sum(v for v in result.detail[result.detail["门店"] == "二店"]["金额"] if v)
    assert rows[-1][1] == "小计" and rows[-1][5] == pytest.approx(detail_amount)
    assert all(r[1] == "二店" for r in rows[:-1])
    assert result.split_count == 4 and result.split_dir is None


def test_split_into_files(cfg):
    cfg.group_by, cfg.split_by, cfg.split_mode = ["商品"], "门店", "file"
    cfg.pivot_column = "门店"
    result = pipeline.run(cfg)
    folder = result.split_dir
    assert folder == result.output_path.parent / "out_按门店拆分"
    assert sorted(p.name for p in folder.iterdir()) == sorted(
        ["一店.xlsx", "二店.xlsx", "三店.xlsx", "四店.xlsx"])
    wb = openpyxl.load_workbook(folder / "三店.xlsx")
    assert wb.sheetnames == ["汇总", "明细"]
    # 展开列与拆分列相同：单店文件用普通汇总，不重复出“三店/合计”两列
    assert [c.value for c in wb["汇总"][1]] == ["商品", "金额(求和)"]
    assert {r[1] for r in wb["明细"].iter_rows(min_row=2, values_only=True)} == {"三店"}
    assert openpyxl.load_workbook(result.output_path).sheetnames == ["汇总", "明细", "问题清单"]


def test_split_missing_column_and_sheet_limit(cfg, monkeypatch):
    cfg.split_by = "不存在"
    result = pipeline.run(cfg, export=False)
    assert any("拆分列“不存在”" in i.message for i in result.issues.issues)
    monkeypatch.setattr(pipeline, "SPLIT_SHEET_LIMIT", 2)
    cfg.split_by = "门店"
    result = pipeline.run(cfg)
    assert any("超过 2 个" in i.message for i in result.issues.issues)
    assert openpyxl.load_workbook(result.output_path).sheetnames == ["汇总", "明细", "问题清单"]


def test_safe_names():
    from core.exporter import safe_sheet_title
    used = {"汇总"}
    assert safe_sheet_title("汇总", used) == "汇总_2"
    assert safe_sheet_title("a/b:c*?[x]", used) == "a_b_c___x_"
    assert len(safe_sheet_title("长" * 40, used)) == 31
    names = set()
    assert pipeline.safe_filename('CON', names) == "_CON"
    assert pipeline.safe_filename('a<b>:"|?.', names) == "a_b_____"
    assert pipeline.safe_filename('a<b>:"|?.', names) == "a_b______2"


def test_pivot_split_config_roundtrip(cfg, tmp_path):
    cfg.pivot_column, cfg.split_by, cfg.split_mode = "门店", "门店", "file"
    assert TaskConfig.load(cfg.save(tmp_path / "c.json")) == cfg
    assert TaskConfig.from_dict({"input_folder": "x"}).split_mode == "sheet"


def test_cli_pivot_split(data_dir, tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    assert cli.main(["-i", str(data_dir), "-g", "商品", "-s", "金额", "--pivot", "门店",
                     "--split", "门店", "--split-mode", "file", "-o", str(out)]) == 0
    assert "拆分文件" in capsys.readouterr().out
    assert (tmp_path / "r_按门店拆分" / "一店.xlsx").exists()
