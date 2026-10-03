"""阶段 3 健壮性：问题清单逐条核对 + 各种异常文件不崩溃。"""
import datetime as dt

import openpyxl
import pytest

from config.task_config import AggSpec, TaskConfig
from core import pipeline, reader
from core.validator import ERROR, IssueCollector

# 全部测试数据（含子文件夹）应产生的问题清单，逐条核对（顺序即问题清单的排序）
EXPECTED_ISSUES = [
    ("错误", "08_文本数字.xlsx", "Sheet1", 7, "金额", "数值无法转换", "约52元"),
    ("错误", "09_日期混用.xlsx", "Sheet1", 10, "日期", "日期无法识别", "2026-13-45"),
    ("错误", "09_日期混用.xlsx", "Sheet1", 11, "日期", "日期无法识别", "下周一"),
    ("错误", "10_占位值.xlsx", "Sheet1", 7, "数量", "数值无法转换", "很多"),
    ("错误", "16_损坏文件.xlsx", "", None, "", "文件读取失败", ""),
    ("警告", "02_表头同义.xlsx", "Sheet1", 1, "单价/元", "表头待确认", ""),
    ("警告", "11_重复记录.xlsx", "Sheet1", 3, "", "重复记录", ""),
    ("警告", "11_重复记录.xlsx", "Sheet1", 6, "", "重复记录", ""),
    ("警告", "11_重复记录.xlsx", "Sheet1", 7, "", "重复记录", ""),
    ("警告", "12_多Sheet.xlsx", "说明", 1, "", "未识别到标准列", ""),
    ("警告", "17_空文件.xlsx", "Sheet1", None, "", "空文件/空Sheet", ""),
    ("警告", "18_空文件.csv", "", None, "", "空文件/空Sheet", ""),
    ("警告", "19_只有表头.xlsx", "Sheet1", 1, "", "无数据行", ""),
    ("提示", "04_多一列.xlsx", "Sheet1", 1, "会员卡号", "表头未匹配", ""),
    ("提示", "05_少一列.xlsx", "Sheet1", 1, "", "缺少列", ""),
    ("提示", "06_标题行和合计行.xlsx", "Sheet1", 4, "", "表头识别", ""),
    ("提示", "06_标题行和合计行.xlsx", "Sheet1", 13, "", "跳过合计行", "合计"),
    ("提示", "06_标题行和合计行.xlsx", "Sheet1", 14, "", "跳过说明行", "制表人：王五"),
]


def test_issue_list_exact(data_dir):
    cfg = TaskConfig(input_folder=str(data_dir), recursive=True, group_by=["门店"],
                     aggregations=[AggSpec("金额")])
    result = pipeline.run(cfg, export=False)
    got = [(i.severity, i.file, i.sheet, i.row, i.column, i.type,
            "" if i.value is None else str(i.value)) for i in result.issues.sorted()]
    assert got == EXPECTED_ISSUES
    assert result.rows_detail == 130


# ---------------------------------------------------------------- 各种异常文件
def _run(folder, tmp_path, **kw):
    cfg = TaskConfig(input_folder=str(folder), output_dir=str(tmp_path / "out"),
                     output_name="r.xlsx", **kw)
    return pipeline.run(cfg)


def _xlsx(path, rows, title="Sheet1"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(path)
    return wb


def test_error_cells_and_formulas(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["门店", "金额"])
    ws.append(["一店", 10])
    ws.append(["二店", "=1/0"])            # 没有缓存值的公式
    ws.append(["三店", "#DIV/0!"])
    ws["B4"].data_type = "e"               # Excel 错误值
    wb.save(d / "a.xlsx")
    result = _run(d, tmp_path)
    assert result.output_path is not None
    errs = {(i.row, str(i.value)): i.message for i in result.issues.issues if i.severity == ERROR}
    assert "Excel 错误值" in errs[(4, "#DIV/0!")]    # 错误值被记录，不当作空值


def test_fake_encrypted_and_wrong_extension(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    # OLE 文件头（加密的 xlsx 实际是这种格式）+ 垃圾数据
    (d / "加密.xlsx").write_bytes(bytes.fromhex("D0CF11E0A1B11AE1") + b"\0" * 600)
    (d / "其实是文本.xls").write_text("门店,金额\n一店,1\n", encoding="utf-8")
    _xlsx(d / "ok.xlsx", [["门店", "金额"], ["一店", 5]])
    result = _run(d, tmp_path, group_by=["门店"], aggregations=[AggSpec("金额")])
    failed = {i.file for i in result.issues.issues if i.type == "文件读取失败"}
    assert "加密.xlsx" in failed
    assert result.rows_detail >= 1 and result.output_path is not None


def test_hidden_and_chart_sheets(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["门店", "金额"])
    ws.append(["一店", 1])
    hidden = wb.create_sheet("隐藏")
    hidden.append(["门店", "金额"])
    hidden.append(["二店", 2])
    hidden.sheet_state = "hidden"
    from openpyxl.chart import BarChart, Reference
    cs = wb.create_chartsheet("图表")
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=2))
    cs.add_chart(chart)
    wb.save(d / "a.xlsx")
    assert reader.list_sheets(d / "a.xlsx") == ["Sheet", "隐藏"]   # 图表 Sheet 不读
    assert _run(d, tmp_path).rows_detail == 2


def test_csv_semicolon_ragged_and_latin(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    (d / "a.csv").write_bytes("Datum;Filiale;Betrag\n2026-01-02;Köln;1,5\n2026-01-03;Bonn\n"
                              "2026-01-04;München;2\n2026-01-05;Düsseldorf;3\n".encode("cp1252"))
    result = _run(d, tmp_path)
    assert result.rows_detail == 4
    assert {"Köln", "München", "Düsseldorf"} <= set(result.detail["Filiale"])


@pytest.mark.parametrize("text", ["客户,金额\n张喆,1\n", "门店,金额\n一店,1\n",
                                  "客户,金额\n张喆,1\n王堃,2\n李镕,3\n"])
def test_gbk_with_rare_chars(text):
    assert reader.decode_csv_bytes(text.encode("gbk")) == (text, "gbk")


def test_weird_values_export(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    wb = _xlsx(d / "a.xlsx", [
        ["日期", "门店", "金额", "备注"],
        [dt.date(2026, 1, 1), "一店", 1, "=HYPERLINK(\"x\")"],
        [dt.datetime(2026, 1, 2, 13, 45), "二店", 2.5, "x" * 40000],
        ["2026-01-03", "  三店  ", 1e14, "<&>\"'"],
    ], title="特殊 Sheet 名 😀")
    wb.active["D2"].data_type = "s"      # 源文件里就是以 = 开头的文本
    wb.save(d / "a.xlsx")
    result = _run(d, tmp_path)
    wb = openpyxl.load_workbook(result.output_path)
    ws = wb["明细"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows[0][3] == '=HYPERLINK("x")' and ws["D2"].data_type == "s"
    assert rows[1][0] == dt.datetime(2026, 1, 2, 13, 45)
    assert len(rows[1][3]) == 32767                   # 超长文本截断到 Excel 上限
    assert rows[2][1] == "三店" and rows[2][2] == 1e14 and rows[2][3] == "<&>\"'"
    assert rows[0][-2] == "特殊 Sheet 名 😀"


def test_pre_1900_date(tmp_path):
    """1900 年以前的日期：读取时报问题；导出时按文本写出，不写成错误的序列号。"""
    d = tmp_path / "in"
    d.mkdir()
    _xlsx(d / "a.xlsx", [["日期", "门店"], [dt.date(1899, 6, 1), "一店"]])
    result = _run(d, tmp_path)
    (issue,) = [i for i in result.issues.issues if i.type == "日期无法识别"]
    assert issue.row == 2 and "1900" in issue.message

    import pandas as pd
    from core import exporter
    detail = pd.DataFrame([[dt.date(1899, 6, 1)], [dt.date(2026, 1, 1)]], columns=["日期"],
                          dtype=object)
    out = exporter.export_result(tmp_path / "e.xlsx", detail, detail, [])
    ws = openpyxl.load_workbook(out)["明细"]
    assert ws["A2"].value == "1899-06-01" and ws["A3"].value == dt.datetime(2026, 1, 1)


def test_duplicate_header_names(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    _xlsx(d / "a.xlsx", [["门店", "金额", "金额"], ["一店", 1, 2]])
    result = _run(d, tmp_path)
    assert "金额_2" in result.detail.columns
    assert any(i.type == "表头映射冲突" for i in result.issues.issues)


def test_numeric_only_sheet_and_bad_header_override(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    _xlsx(d / "nums.xlsx", [[1, 2, 3], [4, 5, 6]])
    _xlsx(d / "ok.xlsx", [["门店", "金额"], ["一店", 1]])
    result = _run(d, tmp_path, header_rows={"ok.xlsx|Sheet1": 99})
    types = {(i.file, i.type) for i in result.issues.issues}
    assert ("nums.xlsx", "表头识别") in types
    assert ("ok.xlsx", "表头识别") in types           # 指定行超出范围，改用自动识别
    assert result.rows_detail == 1


def test_output_not_writable(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    _xlsx(d / "ok.xlsx", [["门店", "金额"], ["一店", 1]])
    (tmp_path / "out" / "r.xlsx").mkdir(parents=True)    # 目标路径被占用（类似文件被 Excel 打开）
    result = _run(d, tmp_path)
    assert result.output_path is None
    assert any("无法写入结果文件" in i.message for i in result.issues.issues)
    assert not list((tmp_path / "out").glob("*.~*"))    # 不残留临时文件


def test_many_columns_and_empty_rows(tmp_path):
    d = tmp_path / "in"
    d.mkdir()
    header = ["门店", "金额"] + [f"字段{i}" for i in range(200)]
    rows = [header, ["一店", 1] + list(range(200)), [None] * 202, [None] * 202,
            ["二店", 2] + list(range(200))]
    _xlsx(d / "wide.xlsx", rows)
    result = _run(d, tmp_path)
    assert result.rows_detail == 2 and len(result.detail.columns) == 202 + 3
    assert list(result.detail["原始行号"]) == [2, 5]


def test_reader_never_raises_on_random_bytes(tmp_path):
    import random
    rng = random.Random(1)
    for ext in (".xlsx", ".xls", ".xlsb", ".csv"):
        p = tmp_path / f"r{ext}"
        p.write_bytes(bytes(rng.randrange(256) for _ in range(2048)))
        issues = IssueCollector()
        refs = reader.discover(tmp_path, False, issues)
        for ref in refs:
            reader.load_sheet(ref, issues)
        p.unlink()
