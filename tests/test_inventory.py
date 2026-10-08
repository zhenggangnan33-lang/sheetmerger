"""多仓库月末盘点场景：每个仓库文件埋一种常见的坑，汇总结果必须与标准答案一致。"""
import collections

import pytest

from config.task_config import AggSpec, TaskConfig
from core import pipeline
from generate_inventory_data import generate


@pytest.fixture(scope="module")
def inv(tmp_path_factory):
    folder = tmp_path_factory.mktemp("inv") / "盘点"
    return folder, generate(folder)


def _run(folder, **kw):
    cfg = TaskConfig(input_folder=str(folder), excluded=["仓库盘点模板.xlsx|*"], **kw)
    return pipeline.run(cfg, export=False)


def _by(result, kind):
    return collections.Counter(i.file for i in result.issues.issues if i.type == kind)


def test_totals_per_warehouse_match_answer(inv):
    folder, exp = inv
    r = _run(folder, group_by=["仓库"],
             aggregations=[AggSpec("账面数量"), AggSpec("实盘数量")])
    got = {row[0]: (row[1], row[2]) for row in r.summary.itertuples(index=False, name=None)}
    for wh, totals in exp["totals"].items():
        assert got[wh] == totals, wh
    assert "示例仓" not in got


def test_codes_merge_across_warehouses(inv):
    folder, exp = inv
    r = _run(folder, group_by=["商品编码"], aggregations=[AggSpec("实盘数量")])
    codes = set(r.summary["商品编码"]) - {"总计"}
    assert codes <= exp["catalog_codes"]                 # 补零、BX 编码统一后不会拆成两行
    assert "BX-001" in codes or "BX-002" in codes


def test_each_trap_is_reported(inv):
    folder, exp = inv
    r = _run(folder, group_by=["仓库"], aggregations=[AggSpec("实盘数量")])
    assert _by(r, "文件读取失败") == {"月湖仓.xlsx": 1}
    assert _by(r, "数值无法转换") == {"东钱湖仓.xlsx": 2}          # 单价是 #N/A、#DIV/0!
    assert _by(r, "编码已补齐")["北仑二号仓.xlsx"] == 1
    assert _by(r, "日期补全年份")["梅山仓.xlsx"] == 1
    empty = [i for i in r.issues.issues if i.type == "空值" and i.file == "宁海仓.xlsx"]
    assert empty and f"有 {exp['empty_real']} 个空值" in empty[0].message
    assert _by(r, "文件名与内容不一致") == {"北仑三号仓.xlsx": 1}
    assert any("另一张表" in i.message for i in r.issues.issues if i.file == "镇海仓.xlsx")
    assert not _by(r, "日期无法识别")                              # 梅山仓 5 种写法都认得


def test_rounds_keep_last(inv):
    """象山仓 8 个商品盘了两轮：以复盘为准，初盘行不计入，并在问题清单中逐行说明。"""
    folder, exp = inv
    r = _run(folder, group_by=["仓库"], aggregations=[AggSpec("实盘数量")],
             dedup_columns=["仓库", "商品编码"])
    assert _by(r, "多轮盘点") == {"象山仓.xlsx": exp["duplicate_rounds"]}
    assert not _by(r, "重复记录")["象山仓.xlsx"]
    xs = r.detail[r.detail["仓库"] == "象山仓"]
    assert set(xs["盘点轮次"]) == {"复盘"}


def test_rounds_kept_in_generic_mode(inv):
    folder, exp = inv
    r = _run(folder, group_by=["仓库"], aggregations=[AggSpec("实盘数量")],
             dedup_columns=["仓库", "商品编码"], report_mode="generic")
    assert _by(r, "重复记录")["象山仓.xlsx"] == exp["duplicate_rounds"]
    assert r.report is None


def test_filename_wins_over_template_value(inv):
    folder, _exp = inv
    r = _run(folder, group_by=["仓库"], aggregations=[AggSpec("实盘数量")])
    rows = r.detail[r.detail["来源文件"] == "北仑三号仓.xlsx"]
    assert set(rows["仓库"]) == {"北仑三号仓"}
    r = _run(folder, group_by=["仓库"], aggregations=[AggSpec("实盘数量")], name_from_file=False)
    rows = r.detail[r.detail["来源文件"] == "北仑三号仓.xlsx"]
    assert set(rows["仓库"]) == {"北仑二号仓"}
    assert _by(r, "文件名与内容不一致") == {"北仑三号仓.xlsx": 1}


def test_inventory_report_sheets(inv, tmp_path):
    """不选分组列也能出盘点报表：总览 / 按仓库 / 按商品编码，数字与标准答案一致。"""
    import openpyxl
    folder, exp = inv
    cfg = TaskConfig(input_folder=str(folder), excluded=["仓库盘点模板.xlsx|*"],
                     output_dir=str(tmp_path), output_name="r.xlsx")
    r = pipeline.run(cfg)
    assert r.output_path and r.report is not None
    wb = openpyxl.load_workbook(r.output_path)
    assert wb.sheetnames == ["总览", "按仓库", "按商品编码", "明细", "问题清单"]
    rows = list(wb["按仓库"].iter_rows(values_only=True))
    assert rows[0] == ("仓库", "商品数", "账面数量", "实盘数量", "实盘金额", "盘盈盘亏数量", "未盘商品数")
    got = {row[0]: (row[2], row[3]) for row in rows[1:]}
    assert got == exp["totals"]
    by = {row[0]: row for row in rows[1:]}
    assert by["宁海仓"][6] == exp["empty_real"]
    assert by["宁海仓"][5] == by["宁海仓"][3] - by["宁海仓"][2]
    overview = dict(wb["总览"].iter_rows(min_row=2, values_only=True))
    assert overview["仓库数"] == len(exp["totals"])
    assert overview["账面数量合计"] == sum(b for b, _ in exp["totals"].values())
    assert overview["未盘商品数（实盘数量为空）"] == exp["empty_real"]
    codes = [row[0] for row in wb["按商品编码"].iter_rows(min_row=2, values_only=True)]
    assert set(codes) <= exp["catalog_codes"] and len(codes) == len(set(codes))


def test_amount_is_quantity_times_price(inv):
    folder, _exp = inv
    r = _run(folder)
    d = r.detail
    assert "实盘金额" in d.columns
    for q, p, a in d[["实盘数量", "单价", "实盘金额"]].itertuples(index=False, name=None):
        if q is None or p is None:
            assert a is None
        else:
            assert a == round(q * p, 2)


def test_english_headers_mapped(inv):
    folder, _exp = inv
    r = _run(folder, group_by=["仓库"], aggregations=[AggSpec("实盘数量")])
    dx = r.detail[r.detail["来源文件"] == "大榭仓.xlsx"]
    assert dx["实盘数量"].notna().all() and dx["商品"].notna().all()
    assert not {"Count Qty", "Book Qty", "Item Name"} & set(r.detail.columns)
