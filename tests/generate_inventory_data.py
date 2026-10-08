"""生成“多仓库月末盘点”测试数据：每个仓库的文件埋一种真实场景里常见的坑。

用法：python tests/generate_inventory_data.py [输出文件夹]（默认 tests/inventory_data）
generate() 返回标准答案：{仓库: (账面数量合计, 实盘数量合计)}、不同商品编码数等，供测试比对。
"""
from __future__ import annotations

import datetime as dt
import random
import shutil
import sys
from pathlib import Path

import openpyxl

DATE = dt.date(2026, 9, 30)
HEADER = ["盘点日期", "仓库", "商品编码", "商品名称", "规格", "单位", "账面数量", "实盘数量", "单价"]
NAMES = ["矿泉水", "可乐", "橙汁", "牛奶", "咖啡豆", "茶叶", "大米", "面粉", "食用油", "酱油",
         "白糖", "巧克力", "土豆", "西兰花", "柠檬", "牛腩", "鸡腿", "鸡蛋", "芝士", "黄油"]
PACK = ["纸杯", "吸管", "餐盒", "外卖袋", "筷子"]


def catalog(rng: random.Random) -> list[tuple[str, str, float]]:
    """商品目录：40 个 6 位数字编码（多数带前导零）+ 5 个包材编码 BX-00x。"""
    codes = rng.sample(range(100, 30000), 40)
    items = [(f"{c:06d}", rng.choice(NAMES) + rng.choice(["", "A", "B", "精选"]),
              round(rng.uniform(10, 300), 2)) for c in codes]
    items += [(f"BX-{i + 1:03d}", PACK[i], round(rng.uniform(5, 50), 2)) for i in range(5)]
    return items


def records(rng, cat, wh, n=20):
    out = []
    for code, name, price in rng.sample(cat, n):
        book = rng.randint(10, 500)
        out.append({"盘点日期": DATE, "仓库": wh, "商品编码": code, "商品名称": name,
                     "规格": "1kg/袋", "单位": "袋", "账面数量": book,
                     "实盘数量": book + rng.choice([0, 0, 0, -1, 1, -3]), "单价": price})
    return out


def save(path: Path, rows: list[list], merges=(), hidden_rows=(), hidden_cols=(), errors=()):
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    for m in merges:
        ws.merge_cells(m)
    for r in hidden_rows:
        ws.row_dimensions[r].hidden = True
    for c in hidden_cols:
        ws.column_dimensions[c].hidden = True
    for ref, err in errors:
        ws[ref] = err
        ws[ref].data_type = "e"
    wb.save(path)


def table(recs, header=HEADER):
    return [header] + [[r[h] for h in HEADER] for r in recs]


def generate(out: Path) -> dict:
    rng = random.Random(930)
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    cat = catalog(rng)
    expect: dict[str, list] = {}

    def count(recs):
        for r in recs:
            book, real = expect.setdefault(r["仓库"], [0, 0])
            expect[r["仓库"]] = [book + r["账面数量"],
                               real + (r["实盘数量"] if r["实盘数量"] is not None else 0)]

    # 北仑一号仓：规范
    recs = records(rng, cat, "北仑一号仓")
    save(out / "北仑一号仓.xlsx", table(recs))
    count(recs)

    # 北仑二号仓：编码存成数字，前导零丢失
    recs = records(rng, cat, "北仑二号仓")
    rows = table(recs)
    for row in rows[1:]:
        if row[2].isdigit():
            row[2] = int(row[2])
    save(out / "北仑二号仓.xlsx", rows)
    count(recs)

    # 镇海仓：上下两张表（常温区 / 冷链区），第二张表列顺序不同
    recs = records(rng, cat, "镇海仓", 24)
    a, b = recs[:14], recs[14:]
    order2 = ["仓库", "商品编码", "商品名称", "实盘数量", "账面数量", "盘点日期", "规格", "单位", "单价"]
    rows = [["常温区"]] + table(a) + [[], [], ["冷链区"], order2] + [[r[h] for h in order2] for r in b]
    save(out / "镇海仓.xlsx", rows)
    count(recs)

    # 江东仓：两级表头（“数量”合并在上，下分“账面”“实盘”）
    recs = records(rng, cat, "江东仓")
    top = ["盘点日期", "仓库", "商品编码", "商品名称", "规格", "单位", "数量", None, "单价"]
    sub = [None, None, None, None, None, None, "账面", "实盘", None]
    rows = [top, sub] + [[r[h] for h in HEADER] for r in recs]
    merges = ["A1:A2", "B1:B2", "C1:C2", "D1:D2", "E1:E2", "F1:F2", "G1:H1", "I1:I2"]
    save(out / "江东仓.xlsx", rows, merges=merges)
    count(recs)

    # 鄞州仓：盘点日期、仓库两列整列合并
    recs = records(rng, cat, "鄞州仓")
    rows = table(recs)
    for row in rows[2:]:
        row[0] = row[1] = None
    n = len(rows)
    save(out / "鄞州仓.xlsx", rows, merges=[f"A2:A{n}", f"B2:B{n}"])
    count(recs)

    # 慈溪仓：实盘数量带单位
    recs = records(rng, cat, "慈溪仓")
    rows = table(recs)
    for i, row in enumerate(rows[1:]):
        row[7] = f"{row[7]}{' ' if i % 2 else ''}{rng.choice(['箱', '瓶', '包'])}"
    save(out / "慈溪仓.xlsx", rows)
    count(recs)

    # 余姚仓：GBK CSV，编码写成 ="000123"
    recs = records(rng, cat, "余姚仓")
    lines = [",".join(HEADER)]
    for r in recs:
        lines.append(f'{r["盘点日期"]:%Y-%m-%d},{r["仓库"]},="{r["商品编码"]}",{r["商品名称"]},'
                     f'{r["规格"]},{r["单位"]},{r["账面数量"]},{r["实盘数量"]},{r["单价"]}')
    (out / "余姚仓.csv").write_bytes(("\r\n".join(lines) + "\r\n").encode("gbk"))
    count(recs)

    # 宁海仓：5 个商品未盘点（实盘数量为空）
    recs = records(rng, cat, "宁海仓")
    for r in recs[:5]:
        r["实盘数量"] = None
    save(out / "宁海仓.xlsx", table(recs))
    count(recs)

    # 奉化仓：包材编码写法不一
    recs = records(rng, cat, "奉化仓")
    rows = table(recs)
    variants = iter(["bx-{}", " BX-{} ", "BX－{}", "Bx-{}", "bx-{}"])
    for row in rows[1:]:
        if row[2].startswith("BX-"):
            row[2] = next(variants, "BX-{}").format(row[2][3:])
    save(out / "奉化仓.xlsx", rows)
    count(recs)

    # 东钱湖仓：2 个单价是 Excel 错误值
    recs = records(rng, cat, "东钱湖仓")
    save(out / "东钱湖仓.xlsx", table(recs), errors=[("I3", "#N/A"), ("I7", "#DIV/0!")])
    count(recs)

    # 杭州湾仓：筛选隐藏了 6 行，另有隐藏列“成本价(内部)”
    recs = records(rng, cat, "杭州湾仓")
    rows = [HEADER + ["成本价(内部)"]] + [[r[h] for h in HEADER] + [round(r["单价"] * 0.6, 2)]
                                        for r in recs]
    save(out / "杭州湾仓.xlsx", rows, hidden_rows=range(4, 10), hidden_cols=["J"])
    count(recs)

    # 梅山仓：盘点日期是文本，5 种写法
    recs = records(rng, cat, "梅山仓")
    rows = table(recs)
    forms = ["2026.09.30", "9.30", "9月30日", "2026年9月30日", "20260930"]
    for i, row in enumerate(rows[1:]):
        row[0] = forms[i % len(forms)]
    save(out / "梅山仓.xlsx", rows)
    count(recs)

    # 大榭仓：英文表头
    recs = records(rng, cat, "大榭仓")
    en = ["Count Date", "Warehouse", "SKU", "Item Name", "Spec", "UOM", "Book Qty", "Count Qty",
          "Unit Price"]
    save(out / "大榭仓.xlsx", table(recs, en))
    count(recs)

    # 保税区仓：批次级明细，同一商品拆成多个批次
    rows = [HEADER + ["批次号"]]
    batch_recs = []
    for code, name, price in rng.sample(cat, 30):
        for k in range(rng.randint(10, 30)):
            book = rng.randint(0, 3)
            r = {"盘点日期": DATE, "仓库": "保税区仓", "商品编码": code, "商品名称": name,
                 "规格": "1kg/袋", "单位": "袋", "账面数量": book, "实盘数量": book, "单价": price}
            batch_recs.append(r)
            rows.append([r[h] for h in HEADER] + [f"B{code}{k:03d}"])
    save(out / "保税区仓.xlsx", rows)
    count(batch_recs)

    # 象山仓：8 个商品盘了两次（初盘 / 复盘）
    recs = records(rng, cat, "象山仓")
    rows = [HEADER + ["盘点轮次"]] + [[r[h] for h in HEADER] + ["复盘"] for r in recs]
    first = [dict(r, 实盘数量=r["实盘数量"] + 5) for r in recs[:8]]
    rows += [[r[h] for h in HEADER] + ["初盘"] for r in first]
    save(out / "象山仓.xlsx", rows)
    count(recs + first)            # 工具不认识“初盘/复盘”这种业务规则，两轮都会计入（并提示重复）

    # 北仑三号仓：“仓库”列套用模板没改，写成了北仑二号仓
    recs = records(rng, cat, "北仑二号仓")
    save(out / "北仑三号仓.xlsx", table(recs))
    count([dict(r, 仓库="北仑三号仓") for r in recs])     # 口径：以文件名为准

    # 模板文件（测试时排除）与损坏文件
    save(out / "仓库盘点模板.xlsx",
         table([{h: v for h, v in zip(HEADER, [DATE, "示例仓", "000000", "示例商品（请删除此行）",
                                               "", "", 0, 0, 0])}]) + [["填写说明：账面数量取自系统"]])
    (out / "月湖仓.xlsx").write_bytes(b"PK\x03\x04" + bytes(300))

    codes = {c for c, _n, _p in cat}
    return {"totals": {k: tuple(v) for k, v in expect.items()}, "catalog_codes": codes,
            "empty_real": 5, "duplicate_rounds": 8}


if __name__ == "__main__":
    target = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).with_name("inventory_data")
    info = generate(target)
    print(f"已生成：{target}，仓库 {len(info['totals'])} 个")
