"""导出 xlsx：汇总 / 明细 / 问题清单 三个 Sheet。

格式：表头加粗带底色、冻结首行、自动列宽、数值列千分位、日期列统一格式。

写出分两步，兼顾格式规范和速度：
1. openpyxl 生成工作簿骨架：样式表、列宽、冻结窗格、表头行，以及一行"样式样本"
   （每种数字格式各一个单元格，用来让 openpyxl 登记样式并得到样式编号）。
2. 按样本行拿到样式编号后，直接流式生成每个 Sheet 的数据行 XML，替换掉样本行。
   单元格写法与 openpyxl write_only 模式一致（数字 / 行内字符串），Excel、WPS 均可打开。
   逐个创建 openpyxl 单元格对象在 50 万行时要 50 多秒，直接写 XML 只需十几秒。
"""
from __future__ import annotations

import datetime as dt
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from typing import Any, Iterable, Iterator

import pandas as pd
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .validator import ERROR, ISSUE_COLUMNS, WARNING, Issue

SHEET_SUMMARY = "汇总"
SHEET_DETAIL = "明细"
SHEET_ISSUES = "问题清单"

EXCEL_MAX_ROWS = 1_048_576
EXCEL_MAX_CELL_CHARS = 32_767
WIDTH_SAMPLE_ROWS = 3000
MIN_WIDTH, MAX_WIDTH = 6, 60
WRITE_BATCH_ROWS = 2000

FMT_INT = "#,##0"
FMT_FLOAT = "#,##0.00"
FMT_DATE = "yyyy-mm-dd"
FMT_DATETIME = "yyyy-mm-dd hh:mm:ss"

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
_THIN = Side(style="thin", color="A6A6A6")
_HEADER_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=False)
_SEVERITY_FONT = {ERROR: Font(color="C00000", bold=True), WARNING: Font(color="C65911")}

# 样式样本行中各单元格的含义（顺序固定）
_SAMPLE_STYLES: list[tuple[str, dict]] = [
    (FMT_INT, {"number_format": FMT_INT}),
    (FMT_FLOAT, {"number_format": FMT_FLOAT}),
    (FMT_DATE, {"number_format": FMT_DATE}),
    (FMT_DATETIME, {"number_format": FMT_DATETIME}),
    (ERROR, {"font": _SEVERITY_FONT[ERROR]}),
    (WARNING, {"font": _SEVERITY_FONT[WARNING]}),
]

_EXCEL_EPOCH = dt.datetime(1899, 12, 30)
_MIN_EXCEL_DATE = dt.date(1900, 3, 1)       # 更早的日期在 Excel 1900 日期系统中有偏差，按文本写出
_XML_ESCAPE = re.compile(r"[&<>]")
_ESCAPES = {"&": "&amp;", "<": "&lt;", ">": "&gt;"}


def default_filename(now: dt.datetime | None = None) -> str:
    now = now or dt.datetime.now()
    return f"汇总结果_{now:%Y%m%d_%H%M%S}.xlsx"


# ---------------------------------------------------------------- 列格式与列宽
def _display_width(v: Any) -> int:
    if v is None:
        return 0
    if isinstance(v, dt.datetime):
        return 19
    if isinstance(v, dt.date):
        return 10
    if isinstance(v, float):
        s = f"{v:,.2f}"
    elif isinstance(v, int) and not isinstance(v, bool):
        s = f"{v:,}"
    else:
        s = str(v)
    s = s.split("\n", 1)[0][:200]
    return sum(2 if unicodedata.east_asian_width(ch) in "WF" else 1 for ch in s)


def _column_format(values: Iterable[Any]) -> str | None:
    """根据列中的值决定数字格式：整数/小数千分位，日期/日期时间统一格式。"""
    has_num = has_float = has_date = has_dt = has_other = False
    for v in values:
        if v is None or (isinstance(v, float) and v != v):
            continue
        if isinstance(v, bool):
            has_other = True
        elif isinstance(v, int):
            has_num = True
        elif isinstance(v, float):
            has_num = True
            has_float = has_float or not v.is_integer()
        elif isinstance(v, dt.datetime):
            has_dt = True
        elif isinstance(v, dt.date):
            has_date = True
        else:
            has_other = True
    if has_num and not (has_date or has_dt):
        return FMT_FLOAT if has_float else FMT_INT
    if (has_date or has_dt) and not has_num and not has_other:
        return FMT_DATETIME if has_dt else FMT_DATE
    return None


def _safe_value(v: Any) -> Any:
    """把 pandas / numpy 类型转成普通 Python 值，清理 Excel 不接受的字符。"""
    if v is None:
        return None
    if isinstance(v, float) and (v != v or math.isinf(v)):
        return None
    if isinstance(v, pd.Timestamp):
        return None if pd.isna(v) else v.to_pydatetime()
    if v is pd.NaT:
        return None
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):   # numpy 标量
        try:
            return _safe_value(v.item())
        except (ValueError, AttributeError):
            pass
    if isinstance(v, str):
        v = ILLEGAL_CHARACTERS_RE.sub("", v)
        if len(v) > EXCEL_MAX_CELL_CHARS:
            v = v[:EXCEL_MAX_CELL_CHARS]
    return v


# ---------------------------------------------------------------- Sheet 描述
@dataclass
class _SheetSpec:
    title: str
    columns: list[str]
    rows: list[list[Any]]
    formats: list[str | None]
    widths: list[float]
    severity_col: int | None = None
    mixed_cols: frozenset = frozenset()


def _make_specs(title: str, columns: list[str], rows: list[list[Any]],
                severity_col: int | None = None,
                plain_columns: Iterable[str] = ()) -> list[_SheetSpec]:
    plain = set(plain_columns)
    formats = [None if name in plain else _column_format(r[i] for r in rows)
               for i, name in enumerate(columns)]
    widths = []
    for i, name in enumerate(columns):
        w = max([_display_width(name)] + [_display_width(r[i]) for r in rows[:WIDTH_SAMPLE_ROWS]])
        widths.append(min(max(w + 2, MIN_WIDTH), MAX_WIDTH))
    mixed = frozenset(i for i, f in enumerate(formats)
                      if f == FMT_FLOAT and any(isinstance(r[i], str) for r in rows))
    per_sheet = EXCEL_MAX_ROWS - 1
    chunks = [rows[i:i + per_sheet] for i in range(0, len(rows), per_sheet)] or [[]]
    return [_SheetSpec(title if n == 1 else f"{title}_{n}", columns, chunk, formats, widths,
                       severity_col, mixed)
            for n, chunk in enumerate(chunks, start=1)]


# ---------------------------------------------------------------- 第 1 步：openpyxl 骨架
def _write_skeleton(path: Path, specs: list[_SheetSpec]) -> None:
    wb = Workbook(write_only=True)
    for spec in specs:
        ws = wb.create_sheet(spec.title)
        for i, w in enumerate(spec.widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        header = []
        for name in spec.columns:
            c = WriteOnlyCell(ws, value=_safe_value(name))
            c.font, c.fill, c.border, c.alignment = (_HEADER_FONT, _HEADER_FILL,
                                                     _HEADER_BORDER, _HEADER_ALIGN)
            c.data_type = "s"
            header.append(c)
        ws.append(header)
        sample = []
        for _key, attrs in _SAMPLE_STYLES:
            c = WriteOnlyCell(ws, value=0)
            for k, v in attrs.items():
                setattr(c, k, v)
            sample.append(c)
        ws.append(sample)
    wb.save(path)


def _sheet_paths(zf: zipfile.ZipFile) -> dict[str, str]:
    """Sheet 名 -> XML 路径。"""
    from xml.etree import ElementTree as ET
    rel_ns = "{http://schemas.openxmlformats.org/package/2006/relationships}"
    r_id = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
    main_ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    rels = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    targets = {r.get("Id"): r.get("Target") for r in rels.iter(f"{rel_ns}Relationship")}
    workbook = ET.fromstring(zf.read("xl/workbook.xml"))
    paths = {}
    for sheet in workbook.iter(f"{main_ns}sheet"):
        target = targets[sheet.get(r_id)]
        paths[sheet.get("name")] = (target.lstrip("/") if target.startswith("/")
                                    else str(PurePosixPath("xl") / target))
    return paths


# ---------------------------------------------------------------- 第 2 步：流式写数据行
def _esc(s: str) -> str:
    return _XML_ESCAPE.sub(lambda m: _ESCAPES[m.group()], s)


def _rows_xml(spec: _SheetSpec, styles: dict[str, str]) -> Iterator[str]:
    letters = [get_column_letter(i + 1) for i in range(len(spec.columns))]
    col_style = [f' s="{styles[f]}"' if f else "" for f in spec.formats]
    date_style = [col_style[i] if f in (FMT_DATE, FMT_DATETIME) else None
                  for i, f in enumerate(spec.formats)]
    num_style = [col_style[i] if f in (FMT_INT, FMT_FLOAT) else "" for i, f in enumerate(spec.formats)]
    # 数字和文字混排的列（如总览的“值”列）里，整数用整数格式：仓库数 16 不应显示成 16.00
    s_int = f' s="{styles[FMT_INT]}"'
    int_style = [s_int if f == FMT_FLOAT and i in spec.mixed_cols else num_style[i]
                 for i, f in enumerate(spec.formats)]
    s_date, s_dt = f' s="{styles[FMT_DATE]}"', f' s="{styles[FMT_DATETIME]}"'
    sev_col = spec.severity_col
    sev_style = {k: f' s="{styles[k]}"' for k in (ERROR, WARNING)}
    epoch, min_date, illegal = _EXCEL_EPOCH, _MIN_EXCEL_DATE, ILLEGAL_CHARACTERS_RE
    max_chars = EXCEL_MAX_CELL_CHARS

    buf: list[str] = []
    for n, row in enumerate(spec.rows, start=2):
        rn = str(n)
        parts = [f'<row r="{rn}">']
        for j, v in enumerate(row):
            if v is None:
                continue
            t = type(v)
            if t is not str and t not in (int, float, bool, dt.date, dt.datetime):
                v = _safe_value(v)
                if v is None:
                    continue
                t = type(v)
            ref = letters[j] + rn
            if t is str:
                if illegal.search(v):
                    v = illegal.sub("", v)
                if len(v) > max_chars:
                    v = v[:max_chars]
                if "&" in v or "<" in v or ">" in v:
                    v = _esc(v)
                space = ' xml:space="preserve"' if v[:1].isspace() or v[-1:].isspace() else ""
                st = sev_style.get(v, "") if j == sev_col else ""
                parts.append(f'<c r="{ref}"{st} t="inlineStr"><is><t{space}>{v}</t></is></c>')
            elif t is int:
                parts.append(f'<c r="{ref}"{int_style[j]} t="n"><v>{v}</v></c>')
            elif t is float:
                if v != v or v in (math.inf, -math.inf):
                    continue
                parts.append(f'<c r="{ref}"{num_style[j]} t="n"><v>{v!r}</v></c>')
            elif t is bool:
                parts.append(f'<c r="{ref}" t="b"><v>{int(v)}</v></c>')
            else:   # 日期 / 日期时间
                d = v.date() if t is dt.datetime else v
                if d < min_date:
                    parts.append(f'<c r="{ref}" t="inlineStr"><is><t>{v.isoformat()}</t></is></c>')
                    continue
                if t is dt.datetime:
                    delta = v - epoch
                    serial = delta.days + (delta.seconds + delta.microseconds / 1e6) / 86400
                    st = date_style[j] or s_dt
                else:
                    serial = (v - epoch.date()).days
                    st = date_style[j] or s_date
                parts.append(f'<c r="{ref}"{st} t="n"><v>{serial!r}</v></c>')
        parts.append("</row>")
        buf.append("".join(parts))
        if len(buf) >= WRITE_BATCH_ROWS:
            yield "".join(buf)
            buf = []
    if buf:
        yield "".join(buf)


def _sample_styles(xml: str) -> tuple[int, int, dict[str, str]]:
    """在骨架 Sheet 中找到样本行，返回 (样本行起点, 样本行终点, 样式编号表)。"""
    start = xml.index('<row r="2"')
    end = xml.index("</row>", start) + len("</row>")
    cells = re.findall(r"<c\b([^>]*)>", xml[start:end])
    styles = {}
    for (key, _attrs), cell in zip(_SAMPLE_STYLES, cells):
        m = re.search(r'\bs="(\d+)"', cell)
        styles[key] = m.group(1) if m else "0"
    return start, end, styles


def _write_final(skeleton: Path, target: Path, specs: list[_SheetSpec]) -> None:
    with zipfile.ZipFile(skeleton) as src, \
            zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as dst:
        sheet_paths = _sheet_paths(src)
        spec_by_path = {sheet_paths[s.title]: s for s in specs}
        for info in src.infolist():
            data = src.read(info.filename)
            spec = spec_by_path.get(info.filename)
            if spec is None:
                dst.writestr(info, data)
                continue
            xml = data.decode("utf-8")
            start, end, styles = _sample_styles(xml)
            zinfo = zipfile.ZipInfo(info.filename, date_time=info.date_time)
            zinfo.compress_type = zipfile.ZIP_DEFLATED
            # ZipInfo 默认不压缩，必须显式指定（压缩级别属性名在 Python 3.13 改过）
            try:
                zinfo.compress_level = 6
            except AttributeError:
                zinfo._compresslevel = 6
            with dst.open(zinfo, "w", force_zip64=True) as out:
                out.write(xml[:start].encode("utf-8"))
                for chunk in _rows_xml(spec, styles):
                    out.write(chunk.encode("utf-8"))
                out.write(xml[end:].encode("utf-8"))


# ---------------------------------------------------------------- 对外入口
def _frame_rows(df: pd.DataFrame) -> list[list[Any]]:
    return [list(r) for r in df.itertuples(index=False, name=None)]


_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
MAX_SHEET_TITLE = 31


def safe_sheet_title(name: str, used: set[str]) -> str:
    """Excel Sheet 名：去掉非法字符、限 31 字、不重名（不区分大小写）。"""
    base = _BAD_SHEET_CHARS.sub("_", str(name)).strip().strip("'") or "Sheet"
    base = base[:MAX_SHEET_TITLE]
    title, n = base, 2
    while title.lower() in used:
        suffix = f"_{n}"
        title = base[:MAX_SHEET_TITLE - len(suffix)] + suffix
        n += 1
    used.add(title.lower())
    return title


def export_result(path: str | Path, summary: pd.DataFrame | None, detail: pd.DataFrame,
                  issues: list[Issue] | None, plain_columns: Iterable[str] = (),
                  extra_sheets: Iterable[tuple[str, list[str], list[list[Any]]]] = ()) -> Path:
    """写出结果文件。

    Sheet 顺序：汇总 → extra_sheets（如按门店拆分的
    各 Sheet）→ 明细 → 问题清单。summary 为 None 时不写汇总 Sheet。
    issues 需已按严重程度排序；传 None 则不写问题清单。plain_columns 中的列（如行号）不加千分位。
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    plain_columns = list(plain_columns)
    used = {SHEET_SUMMARY.lower(), SHEET_DETAIL.lower(), SHEET_ISSUES.lower()}
    specs: list[_SheetSpec] = []
    if summary is not None:
        specs += _make_specs(SHEET_SUMMARY, list(summary.columns), _frame_rows(summary))
    for title, columns, rows in extra_sheets:
        specs += _make_specs(safe_sheet_title(title, used), columns, rows,
                             plain_columns=plain_columns)
    specs += _make_specs(SHEET_DETAIL, list(detail.columns), _frame_rows(detail),
                         plain_columns=plain_columns)
    if issues is not None:
        specs += _make_specs(SHEET_ISSUES, ISSUE_COLUMNS, [i.to_row() for i in issues],
                             severity_col=0, plain_columns={"行号"})
    skeleton = path.with_name(path.stem + ".~skeleton.xlsx")
    tmp = path.with_name(path.stem + ".~tmp.xlsx")
    try:
        _write_skeleton(skeleton, specs)
        _write_final(skeleton, tmp, specs)
        tmp.replace(path)
    finally:
        for p in (skeleton, tmp):
            try:
                p.unlink()
            except OSError:
                pass
    return path
